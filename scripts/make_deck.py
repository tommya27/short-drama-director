"""生成 5 页中英双语 PPT（BAYTECH 2026 提交用）。

用法：py -3.13 scripts/make_deck.py
输出：docs/PITCH_DECK.pptx
封面右侧嵌入本项目**真实生成**的场景背板（来自 data/assets/index.json），不用外部素材。
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "PITCH_DECK.pptx"
FONT = "微软雅黑"
INK = RGBColor(0x23, 0x2B, 0x27)
MUTED = RGBColor(0x6B, 0x77, 0x70)
ACCENT = RGBColor(0x2F, 0x6F, 0x5B)


def set_font(run, size: int, bold: bool = False, color: RGBColor = INK, font: str = FONT) -> None:
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font
    # 中文字形需要显式设置 East Asian 字体，否则可能回退成宋体
    rpr = run._r.get_or_add_rPr()
    from pptx.oxml.ns import qn
    ea = rpr.find(qn("a:ea"))
    if ea is None:
        ea = rpr.makeelement(qn("a:ea"), {})
        rpr.append(ea)
    ea.set("typeface", font)


def textbox(slide, left, top, width, height, lines, size=14, spacing=8):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.word_wrap = True
    first = True
    for line in lines:
        para = frame.paragraphs[0] if first else frame.add_paragraph()
        first = False
        para.space_after = Pt(spacing)
        text = line if isinstance(line, str) else line[0]
        style = {} if isinstance(line, str) else line[1]
        run = para.add_run()
        run.text = text
        set_font(run, style.get("size", size), style.get("bold", False), style.get("color", INK))
    return box


def slide_blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def header(slide, cn: str, en: str, index: str):
    textbox(slide, 0.6, 0.35, 12.2, 0.5, [[f"{index}  {cn}", {"size": 24, "bold": True, "color": ACCENT}]])
    textbox(slide, 0.6, 0.95, 12.2, 0.4, [[en, {"size": 13, "color": MUTED}]])


def pick_backdrop() -> Path | None:
    """取本场景最近一次生成的背板作为封面图（真实产物，不是素材库图片）。"""
    index = ROOT / "data" / "assets" / "index.json"
    if not index.exists():
        return None
    try:
        rows = json.loads(index.read_text(encoding="utf-8"))
    except json.JSONDecodeError:
        return None
    for row in sorted(rows, key=lambda r: r.get("created_at", ""), reverse=True):
        if row.get("kind") == "backdrop":
            path = ROOT / "data" / "assets" / str(row.get("file", ""))
            if path.exists():
                return path
    return None


def main() -> int:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # ---------- 1 封面 ----------
    s = slide_blank(prs)
    textbox(s, 0.8, 1.5, 7.2, 1.2, [["可视化 AI 导演台", {"size": 40, "bold": True}]])
    textbox(s, 0.8, 2.6, 7.2, 0.6, [["让角色先演一遍，你来定稿", {"size": 20, "color": ACCENT}]])
    textbox(s, 0.8, 3.2, 7.2, 1.0, [["A Visual Director's Desk for Short Drama", {"size": 15, "color": MUTED}],
                                    ["Let the cast play the scene out — you decide.", {"size": 15, "color": MUTED}]])
    textbox(s, 0.8, 4.5, 7.2, 1.6, [
        ["多个角色在同一沙盘里行动、交流、互相影响；", {"size": 14}],
        ["你观察、干预、确认，再整理成剧本与分镜草稿。", {"size": 14}],
        ["BAYTECH 2026 ｜ AI 应用与工程赛道 ｜ 深圳河套科创中心", {"size": 12, "color": MUTED}],
    ])
    backdrop = pick_backdrop()
    if backdrop:
        s.shapes.add_picture(str(backdrop), Inches(8.4), Inches(1.7), height=Inches(4.2))
        textbox(s, 8.4, 6.0, 4.2, 0.4, [["封面图由本工具现场文生图产出", {"size": 10, "color": MUTED}]])

    # ---------- 2 问题与场景价值 ----------
    s = slide_blank(prs)
    header(s, "问题与场景价值", "The problem, and why it matters", "01")
    textbox(s, 0.7, 1.6, 6.0, 5.0, [
        ["现有 AI 剧本工具给你一段文字，", {"size": 15, "bold": True}],
        ["但你看不见角色为什么这么说 —— 过程是黑盒。", {"size": 15, "bold": True}],
        ["", {"size": 6}],
        ["• 连续性最贵：谁知道什么、道具从哪来、第几集改了什么", {"size": 13}],
        ["• 试戏成本高：只想试一小段，却要重跑整段", {"size": 13}],
        ["• 改好的台词，下次生成就被覆盖", {"size": 13}],
        ["", {"size": 6}],
        ["目标用户：短剧编剧、导演/制片、剧情号创作者、小型编剧工作室", {"size": 12, "color": MUTED}],
    ])
    textbox(s, 7.0, 1.6, 5.8, 5.0, [
        ["AI hands you a scene, but not why.", {"size": 14, "bold": True}],
        ["The process is a black box.", {"size": 14, "bold": True}],
        ["", {"size": 6}],
        ["• Continuity is the expensive part: who knows what, and when", {"size": 12}],
        ["• Rehearsing costs a full rerun, not a single line", {"size": 12}],
        ["• Your edits get overwritten on the next generation", {"size": 12}],
        ["", {"size": 6}],
        ["Users: short-drama writers, directors/producers, scripted short-video creators", {"size": 11, "color": MUTED}],
    ])

    # ---------- 3 24h 增量 ----------
    s = slide_blank(prs)
    header(s, "24 小时增量：我们做了什么", "What we built in 24 hours", "02")
    textbox(s, 0.7, 1.6, 6.0, 5.2, [
        ["1  看得见 Agent 怎么演", {"size": 15, "bold": True, "color": ACCENT}],
        ["     2.5D 状态地图 + three.js 舞台；点事件看", {"size": 12}],
        ["     提议 → 检查 → 最终动作 → 状态变化", {"size": 12}],
        ["2  剧情事实透镜", {"size": 15, "bold": True, "color": ACCENT}],
        ["     谁持有、谁已知、依据是什么；非法动作", {"size": 12}],
        ["     可见降级并记录原因，不静默改写、不毁整轮", {"size": 12}],
        ["3  作者掌控", {"size": 15, "bold": True, "color": ACCENT}],
        ["     候选→预览→提交（幂等+版本校验）；字段锁定、", {"size": 12}],
        ["     导演指令、分支试演；没有自动提交路径", {"size": 12}],
        ["4  真模型 + 真素材", {"size": 15, "bold": True, "color": ACCENT}],
        ["     角色候选由真模型生成（source=llm）；背板与", {"size": 12}],
        ["     立绘由文生图产出，带提示词/模型/时间戳", {"size": 12}],
        ["5  一句话生成完整开场", {"size": 15, "bold": True, "color": ACCENT}],
        ["     模型产出剧名/区域/角色/道具/事实 + 舞台布局；", {"size": 12}],
        ["     华山→室外山景，游轮→驾驶舱，来源如实标注", {"size": 12}],
    ])
    textbox(s, 7.0, 1.6, 5.8, 5.2, [
        ["1  See how agents act", {"size": 13, "bold": True, "color": ACCENT}],
        ["     2.5D map + three.js stage; click to see", {"size": 11}],
        ["     proposal → fact check → final action", {"size": 11}],
        ["2  Fact lens", {"size": 13, "bold": True, "color": ACCENT}],
        ["     who holds / who knows / on what evidence;", {"size": 11}],
        ["     invalid moves degrade visibly, never silently", {"size": 11}],
        ["3  Author stays in control", {"size": 13, "bold": True, "color": ACCENT}],
        ["     draft → preview → commit (idempotent,", {"size": 11}],
        ["     revision-checked); locks, directives, branches", {"size": 11}],
        ["4  Real model, real assets", {"size": 13, "bold": True, "color": ACCENT}],
        ["     live LLM candidates (source=llm); text-to-image", {"size": 11}],
        ["     backdrops and portraits with provenance", {"size": 11}],
        ["5  One line to a full opening", {"size": 13, "bold": True, "color": ACCENT}],
        ["     title, areas, cast, props, facts + stage layout;", {"size": 11}],
        ["     mountain vs ship bridge, source labelled", {"size": 11}],
    ])

    # ---------- 4 技术与边界 ----------
    s = slide_blank(prs)
    header(s, "技术与边界（诚实版）", "Stack and honest boundaries", "03")
    textbox(s, 0.7, 1.6, 6.0, 5.2, [
        ["架构", {"size": 15, "bold": True, "color": ACCENT}],
        ["React + three.js 导演台 ｜ FastAPI 剧情内核（SQLite 版本化）", {"size": 12}],
        ["｜ 真实模型（OpenAI 兼容）+ 文生图（Seedream）", {"size": 12}],
        ["可切换运行模式", {"size": 15, "bold": True, "color": ACCENT}],
        ["live（真模型）/ offline（本地确定性规则）；模式徽标常驻；", {"size": 12}],
        ["显式 live 缺密钥会拒绝启动，绝不静默降级", {"size": 12}],
        ["复用边界", {"size": 15, "bold": True, "color": ACCENT}],
        ["叙事内核迁移自用户自己的旧工作区，记录源文件 sha256 与改写范围；", {"size": 12}],
        ["3D 布局参考 FeiControl（MIT），未复制其代码与资产", {"size": 12}],
        ["不承诺", {"size": 15, "bold": True, "color": ACCENT}],
        ["不是一键生成整部作品；3D 未达电影级；事实判定可能漏判误判", {"size": 12}],
    ])
    textbox(s, 7.0, 1.6, 5.8, 5.2, [
        ["Stack", {"size": 13, "bold": True, "color": ACCENT}],
        ["React + three.js console; FastAPI narrative core with", {"size": 11}],
        ["versioned SQLite; live LLM (OpenAI-compatible) + text-to-image", {"size": 11}],
        ["Switchable runtime", {"size": 13, "bold": True, "color": ACCENT}],
        ["live vs offline, with a persistent badge; explicit live", {"size": 11}],
        ["without a key refuses to start — never a silent downgrade", {"size": 11}],
        ["Reuse boundary", {"size": 13, "bold": True, "color": ACCENT}],
        ["core ported from the user's own earlier workspace (sha256 recorded);", {"size": 11}],
        ["3D layout references FeiControl (MIT) without copying code/assets", {"size": 11}],
        ["Not claimed", {"size": 13, "bold": True, "color": ACCENT}],
        ["not one-click full-length generation; not cinematic 3D;", {"size": 11}],
        ["fact checking is a graded advisory, not infallible", {"size": 11}],
    ])

    # ---------- 5 演示与下一步 ----------
    s = slide_blank(prs)
    header(s, "现场演示与下一步", "Live demo, and what's next", "04")
    textbox(s, 0.7, 1.6, 6.0, 5.2, [
        ["现场真跑", {"size": 15, "bold": True, "color": ACCENT}],
        ["输一句想法 → 模型生成完整开场 → 确认建场景", {"size": 12}],
        ["推进一步 → 点事件看「提议→检查→结果」", {"size": 12}],
        ["→ 点道具看事实 → 改一句台词并锁定 → 加导演指令", {"size": 12}],
        ["→ 预览（正式状态不变）→ 提交 → 生成美术素材", {"size": 12}],
        ["→ 导出场次卡 / 剧本 / 分镜草稿", {"size": 12}],
        ["下一步", {"size": 15, "bold": True, "color": ACCENT}],
        ["真模型长时稳定性与费用评估；image→3D 资产；", {"size": 12}],
        ["更多场景模板与多人协作", {"size": 12}],
        ["", {"size": 8}],
        ["一句话：把最难观察和修改的角色互动，", {"size": 14, "bold": True, "color": ACCENT}],
        ["变成看得见、能干预、可追溯的工作台。", {"size": 14, "bold": True, "color": ACCENT}],
    ])
    textbox(s, 7.0, 1.6, 5.8, 5.2, [
        ["Live demo", {"size": 13, "bold": True, "color": ACCENT}],
        ["create scene → step → inspect an event (proposal → check → result)", {"size": 11}],
        ["→ inspect a prop's facts → edit and lock a line → add a directive", {"size": 11}],
        ["→ preview (no state change) → commit → generate art assets", {"size": 11}],
        ["→ export scene card / script / storyboard", {"size": 11}],
        ["Next", {"size": 13, "bold": True, "color": ACCENT}],
        ["long-run stability and cost of live mode; image-to-3D assets;", {"size": 11}],
        ["more scene templates and collaboration", {"size": 11}],
        ["", {"size": 8}],
        ["Make the hardest part of drama — how characters react to", {"size": 12, "bold": True, "color": ACCENT}],
        ["each other — visible, steerable and traceable.", {"size": 12, "bold": True, "color": ACCENT}],
    ])

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT))
    print(f"已生成 {OUT} （{len(prs.slides.__iter__.__self__._sldIdLst)} 页，封面图={'有' if backdrop else '无'}）")
    return 0


if __name__ == "__main__":
    sys.exit(main())
